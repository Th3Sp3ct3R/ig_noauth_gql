#!/usr/bin/env python3
"""dashboard.py - local web dashboard for ig_noauth_scrape.

    python dashboard.py            # opens http://127.0.0.1:8765 in your browser

Runs entirely on this machine. Load a username list, tick the fields you want, manage a pool
of proxies, start, watch progress, preview and download results (Excel / CSV).
State on disk, next to this file:
    dashboard_results.jsonl   scraped rows (resume source)
    dashboard_found.txt       usernames harvested from posts
    dashboard_proxies.json    proxy pool
"""
import csv, io, json, os, queue, re, sys, threading, time, webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

import ig_noauth_scrape as sc
from ig_gql import IGWeb, Blocked

HOST, PORT = '127.0.0.1', 8765
STORE, FOUND, PROXIES = 'dashboard_results.jsonl', 'dashboard_found.txt', 'dashboard_proxies.json'
BLOCK_COOLDOWN = 3 * 3600   # seconds a blocked proxy / direct IP is rested before reuse

FIELD_GROUPS = [
    ('bio', 'Bio', ['bio']), ('email', 'Email', ['email']), ('phone', 'Phone', ['phone']),
    ('followers', 'Followers', ['followers']), ('following', 'Following', ['following']),
    ('posts', 'Post count (Instagram hides it logged-out)', ['posts']),
    ('likes', 'Likes of last N posts', None), ('dates', 'Dates of last N posts', None),
    ('name', 'Full name', ['full_name']), ('flags', 'Verified / private', ['verified', 'private']),
    ('link', 'External link', ['external_url']),
]


# ---------------------------------------------------------------- proxies
class ProxyPool:
    """Entries: {id, url, rotate, enabled, status, note, blocked_at}. id 'direct' = this computer's IP."""
    def __init__(self):
        self.lock = threading.Lock()
        self.items = []
        self.in_use = {}   # proxy id -> number of threads currently using it
        if os.path.exists(PROXIES):
            try: self.items = json.load(open(PROXIES, encoding='utf-8'))
            except ValueError: self.items = []
        if not any(p['id'] == 'direct' for p in self.items):
            self.items.insert(0, {'id': 'direct', 'url': '', 'rotate': '', 'enabled': True,
                                  'status': 'untested', 'note': '', 'blocked_at': 0})
        self.save()

    def save(self):
        json.dump(self.items, open(PROXIES, 'w', encoding='utf-8'), indent=1)

    def list(self):
        with self.lock:
            now = time.time()
            out = []
            for p in self.items:
                q = dict(p)
                if p['status'] == 'blocked' and now - p['blocked_at'] > BLOCK_COOLDOWN:
                    q['status'] = 'rested'
                q['resting_left'] = max(0, int(BLOCK_COOLDOWN - (now - p['blocked_at']))) if p['status'] == 'blocked' else 0
                out.append(q)
            return out

    def add(self, url, rotate):
        with self.lock:
            pid = 'p%d' % int(time.time() * 1000)
            self.items.append({'id': pid, 'url': url.strip(), 'rotate': rotate.strip(), 'enabled': True,
                               'status': 'untested', 'note': '', 'blocked_at': 0})
            self.save(); return pid

    def add_many(self, lines):
        """Bulk import. One proxy per line: 'url' or 'url,rotate_link' or 'url rotate_link'.
        Skips blanks, comment lines starting with #, and urls already in the pool."""
        existing = {p['url'] for p in self.items}
        added = skipped = 0
        with self.lock:
            ms = int(time.time() * 1000)
            for raw in lines:
                line = raw.strip()
                if not line or line.startswith('#'):
                    continue
                parts = [x.strip() for x in re.split(r'[,	]| {2,}| (?=https?://)', line, maxsplit=1) if x.strip()]
                url = parts[0]
                rotate = parts[1] if len(parts) > 1 else ''
                if '://' not in url:
                    url = 'http://' + url
                if url in existing:
                    skipped += 1; continue
                ms += 1
                self.items.append({'id': f'p{ms}', 'url': url, 'rotate': rotate, 'enabled': True,
                                   'status': 'untested', 'note': '', 'blocked_at': 0})
                existing.add(url); added += 1
            self.save()
        return added, skipped

    def remove(self, pid):
        with self.lock:
            self.items = [p for p in self.items if p['id'] != pid or pid == 'direct']; self.save()

    def remove_many(self, pids):
        ids = set(pids) - {'direct'}
        with self.lock:
            before = len(self.items)
            self.items = [p for p in self.items if p['id'] not in ids]
            self.save(); return before - len(self.items)

    def toggle(self, pid, enabled):
        with self.lock:
            for p in self.items:
                if p['id'] == pid: p['enabled'] = bool(enabled)
            self.save()

    def set_status(self, pid, status, note=''):
        with self.lock:
            for p in self.items:
                if p['id'] == pid:
                    p['status'] = status; p['note'] = note
                    p['blocked_at'] = time.time() if status == 'blocked' else p.get('blocked_at', 0)
            self.save()

    def available(self):
        now = time.time()
        with self.lock:
            return [dict(p) for p in self.items if p['enabled'] and p['status'] != 'failed' and
                    not (p['status'] == 'blocked' and now - p['blocked_at'] < BLOCK_COOLDOWN)]

    def claim(self):
        """Hand a usable connection to a thread, spreading threads across distinct connections
        first and only doubling up once every usable connection already has a claimant."""
        now = time.time()
        with self.lock:
            cands = [p for p in self.items if p['enabled'] and p['status'] != 'failed' and
                     not (p['status'] == 'blocked' and now - p['blocked_at'] < BLOCK_COOLDOWN)]
            if not cands:
                return None
            cands.sort(key=lambda p: self.in_use.get(p['id'], 0))
            chosen = dict(cands[0])
            self.in_use[chosen['id']] = self.in_use.get(chosen['id'], 0) + 1
            return chosen

    def release(self, pid):
        with self.lock:
            if pid in self.in_use:
                self.in_use[pid] = max(0, self.in_use[pid] - 1)

    def test(self, pid):
        p = next((x for x in self.list() if x['id'] == pid), None)
        if not p: return {'ok': False, 'note': 'unknown proxy'}
        t = time.time()
        try:
            ig = IGWeb(proxy=p['url'] or None).warmup()
            r = ig.resolve('instagram', posts=1)
            ms = int((time.time() - t) * 1000)
            if r and r.get('pk'):
                self.set_status(pid, 'ok', f'{ms} ms'); return {'ok': True, 'note': f'{ms} ms'}
            self.set_status(pid, 'failed', 'unexpected answer'); return {'ok': False, 'note': 'unexpected answer'}
        except Blocked as e:
            self.set_status(pid, 'blocked', str(e)); return {'ok': False, 'note': f'blocked ({e})'}
        except Exception as e:
            note = str(e)[:120]
            self.set_status(pid, 'failed', note); return {'ok': False, 'note': note}


POOL = ProxyPool()


class TestAll:
    """Runs Test on a batch of proxies in parallel threads so 'Test all' doesn't block on 500 of them."""
    def __init__(self):
        self.lock = threading.Lock()
        self.stop = threading.Event()
        self.thread = None
        self.total = 0
        self.done = 0
        self.ok = 0

    def running(self):
        return self.thread is not None and self.thread.is_alive()

    def start(self, ids):
        if self.running():
            return None
        self.stop.clear(); self.total, self.done, self.ok = len(ids), 0, 0
        self.thread = threading.Thread(target=self._run, args=(ids,), daemon=True)
        self.thread.start()
        return self.thread

    def _run(self, ids):
        q = queue.Queue()
        for i in ids:
            q.put(i)

        def worker():
            while not self.stop.is_set():
                try:
                    pid = q.get_nowait()
                except queue.Empty:
                    return
                r = POOL.test(pid)
                with self.lock:
                    self.done += 1
                    self.ok += 1 if r.get('ok') else 0
        ths = [threading.Thread(target=worker, daemon=True) for _ in range(min(20, max(1, len(ids))))]
        [t.start() for t in ths]
        [t.join() for t in ths]

    def status(self):
        with self.lock:
            return {'running': self.running(), 'total': self.total, 'done': self.done, 'ok': self.ok}


TESTALL = TestAll()


# ---------------------------------------------------------------- job
class Job:
    def __init__(self):
        self.lock = threading.Lock()
        self.stop = threading.Event()
        self.thread = None
        self.results, self.order, self.found, self.log = {}, [], set(), []
        self.cfg, self.t0, self.t1, self.message, self.proxy_now = {}, None, None, '', ''
        self.tlock = threading.Lock()
        self.thread_state = {}   # thread index -> {proxy, action, done}
        self.load()

    def load(self):
        if os.path.exists(STORE):
            for line in open(STORE, encoding='utf-8'):
                try: self._put(json.loads(line), persist=False)
                except ValueError: pass
        if os.path.exists(FOUND):
            self.found = {l.strip() for l in open(FOUND, encoding='utf-8') if l.strip()}

    def _put(self, row, persist=True):
        u = row['username']
        if u not in self.results: self.order.append(u)
        self.results[u] = row
        if persist:
            with open(STORE, 'a', encoding='utf-8') as f: f.write(json.dumps(row, ensure_ascii=False) + '\n')

    def clear(self):
        with self.lock:
            self.results.clear(); self.order.clear(); self.found.clear(); self.log.clear()
            for p in (STORE, FOUND):
                if os.path.exists(p): os.remove(p)
            self.message = 'Results cleared.'; self.cfg = {}

    def running(self):
        return self.thread is not None and self.thread.is_alive()

    def start(self, cfg):
        if self.running(): return False
        self.cfg, self.message, self.log = cfg, '', []
        self.stop.clear(); self.t0, self.t1 = time.time(), None
        self.thread = threading.Thread(target=self._run, daemon=True); self.thread.start()
        return True

    def _say(self, s):
        with self.lock:
            self.log.append(time.strftime('%H:%M:%S ') + s); self.log = self.log[-500:]

    def _run(self):
        cfg = self.cfg
        sc.NPOSTS = int(cfg.get('nposts', 5))
        names = cfg['usernames']
        done = {u for u, r in self.results.items() if r['status'] != 'blocked' and not r['status'].startswith('error')}
        todo = [u for u in names if u not in done]
        workers = max(1, int(cfg.get('workers', 4)))
        per_ip = int(cfg.get('per_ip', 140))
        harvest = bool(cfg.get('discover'))
        known = set(names) | self.found
        self._say(f'{len(todo)} to do, {len(names) - len(todo)} already done, {workers} thread(s)')

        q = queue.Queue()
        for u in todo:
            q.put(u)
        retries = {}   # username -> failed attempts this run, so one bad account can't loop forever
        ERROR_RETRIES = 2

        with self.tlock:
            self.thread_state.clear()

        def set_state(idx, **kw):
            with self.tlock:
                st = self.thread_state.setdefault(idx, {'proxy': '', 'action': '', 'done': 0})
                st.update(kw)

        def bump_done(idx):
            with self.tlock:
                self.thread_state.setdefault(idx, {'proxy': '', 'action': '', 'done': 0})['done'] += 1

        def clear_state(idx):
            with self.tlock:
                self.thread_state.pop(idx, None)

        def write(row):
            found = row.pop('_found', set())
            with self.lock:
                self._put(row)
                if harvest:
                    new = sorted(h for h in found if h not in known)
                    if new:
                        known.update(new); self.found.update(new)
                        with open(FOUND, 'a', encoding='utf-8') as f: f.write('\n'.join(new) + '\n')
            self._say(f'{row["username"]}: {row["status"]}')

        def label(p): return 'direct connection' if p['id'] == 'direct' else p['url'].split('@')[-1]

        def rotate(p):
            if p.get('rotate'):
                try:
                    import requests; requests.get(p['rotate'], timeout=40); self._say(f'{label(p)}: asked for a new IP')
                except Exception as e:
                    self._say(f'{label(p)}: rotate failed: {str(e)[:90]}')
                time.sleep(3)

        def worker(idx):
            my_proxy, ig, n_on_proxy = None, None, 0
            try:
                while not self.stop.is_set():
                    if q.empty():
                        return
                    if my_proxy is None:
                        set_state(idx, proxy='', action='waiting for a connection')
                        p = POOL.claim()
                        if p is None:
                            time.sleep(1)
                            continue
                        my_proxy, n_on_proxy = p, 0
                        set_state(idx, proxy=label(p), action='connecting')
                        try:
                            ig = IGWeb(proxy=p['url'] or None).warmup()
                        except Blocked as e:
                            POOL.set_status(p['id'], 'blocked', str(e)); POOL.release(p['id'])
                            self._say(f'{label(p)}: blocked at warmup, thread {idx + 1} trying another connection')
                            my_proxy = None; continue
                        except Exception as e:
                            POOL.set_status(p['id'], 'failed', str(e)[:120]); POOL.release(p['id'])
                            self._say(f'{label(p)}: failed ({str(e)[:80]}) - disabled until Tested again')
                            my_proxy = None; continue
                        self._say(f'thread {idx + 1}: using {label(p)}')
                    try:
                        u = q.get_nowait()
                    except queue.Empty:
                        POOL.release(my_proxy['id']); my_proxy = None
                        clear_state(idx)
                        return
                    set_state(idx, action=f'scraping {u}')
                    try:
                        rec = sc.scrape_one(ig, u)
                        write(rec)
                        bump_done(idx)
                        n_on_proxy += 1
                        if per_ip and n_on_proxy >= per_ip:
                            POOL.set_status(my_proxy['id'], 'ok', 'working')
                            self._say(f'thread {idx + 1}: {label(my_proxy)} reached {per_ip} accounts, rotating')
                            set_state(idx, action='rotating connection')
                            rotate(my_proxy); POOL.release(my_proxy['id'])
                            my_proxy, ig = None, None
                    except Blocked:
                        q.put(u)
                        POOL.set_status(my_proxy['id'], 'blocked', 'rate limited'); POOL.release(my_proxy['id'])
                        self._say(f'{label(my_proxy)}: rate limited, thread {idx + 1} switching connection')
                        my_proxy, ig = None, None
                    except Exception as e:
                        n = retries[u] = retries.get(u, 0) + 1
                        if n <= ERROR_RETRIES:
                            q.put(u)
                            self._say(f'{label(my_proxy)}: error on {u} ({str(e)[:70]}), '
                                      f'retry {n}/{ERROR_RETRIES} on another connection')
                            set_state(idx, action='switching connection after an error')
                            POOL.release(my_proxy['id']); my_proxy, ig = None, None
                        else:
                            write({'username': u, 'status': 'error: ' + str(e)[:80]})
            finally:
                if my_proxy is not None:
                    POOL.release(my_proxy['id'])
                clear_state(idx)

        try:
            ths = [threading.Thread(target=worker, args=(i,), daemon=True) for i in range(workers)]
            [t.start() for t in ths]
            stall_ticks = 0
            while any(t.is_alive() for t in ths):
                labs = sorted({s['proxy'] for s in self.thread_status() if s['proxy']})
                self.proxy_now = ', '.join(labs) if labs else ('waiting for a connection' if not q.empty() else '')
                if not labs and not q.empty() and not self.stop.is_set():
                    stall_ticks = stall_ticks + 1 if not POOL.available() else 0
                    if stall_ticks >= 3:
                        nxt = min((p for p in POOL.list() if p['enabled'] and p['status'] == 'blocked'),
                                  key=lambda p: p['resting_left'], default=None)
                        when = f' Earliest one frees up in {nxt["resting_left"] // 60} min.' if nxt else ''
                        self.message = (f'All connections are blocked or disabled. {q.qsize()} usernames left.'
                                        f'{when} Add a proxy on the Proxies page or wait, then press Start again.')
                        self._say(self.message); self.stop.set()
                else:
                    stall_ticks = 0
                time.sleep(1)
            [t.join(timeout=2) for t in ths]
            left = q.qsize()
            if self.stop.is_set() and left:
                self.message = self.message or f'Stopped with {left} usernames left. Press Start to continue.'
            elif not left:
                self.message = 'Finished.'
        except Exception as e:
            self.message = f'Run failed: {e}'; self._say(self.message)
        with self.tlock:
            self.thread_state.clear()
        self.proxy_now = ''; self.t1 = time.time()

    def thread_status(self):
        with self.tlock:
            return [{'thread': i + 1, 'proxy': s.get('proxy', ''), 'action': s.get('action', ''), 'done': s.get('done', 0)}
                    for i, s in sorted(self.thread_state.items())]

    def status(self):
        with self.lock:
            rows = list(self.results.values()); names = self.cfg.get('usernames', [])
            st = {'ok': 0, 'not_found': 0, 'no_posts_visible': 0, 'blocked': 0, 'error': 0}
            for r in rows:
                k = 'error' if r['status'].startswith('error') else r['status']; st[k] = st.get(k, 0) + 1
            done = sum(1 for u in names if u in self.results and self.results[u]['status'] != 'blocked'
                       and not self.results[u]['status'].startswith('error'))
            elapsed = ((self.t1 or time.time()) - self.t0) if self.t0 else 0
            return {'running': self.running(), 'total': len(names), 'done': done, 'stored': len(rows),
                    'counts': st, 'found': len(self.found), 'elapsed': round(elapsed), 'message': self.message,
                    'log': self.log[-80:], 'proxy_now': self.proxy_now, 'threads': self.thread_status()}

    def recent(self, n=60):
        with self.lock:
            return [self.results[u] for u in self.order[-n:]][::-1]

    # ---- export
    def columns(self, fields, nposts):
        cols = ['username', 'status']
        for key, _, produced in FIELD_GROUPS:
            if key not in fields: continue
            if key == 'likes': cols += [f'likes_post{i}' for i in range(1, nposts + 1)]
            elif key == 'dates': cols += [f'date_post{i}' for i in range(1, nposts + 1)]
            else: cols += produced
        return cols

    @staticmethod
    def headers(cols):
        h = dict(sc.HEADERS)
        h.update({f'likes_post{i}': f'Likes post {i}' for i in range(1, 51)})
        h.update({f'date_post{i}': f'Date post {i}' for i in range(1, 51)})
        return [h.get(c, c) for c in cols]

    def export(self, fmt, fields, nposts):
        cols = self.columns(fields, nposts); rows = [self.results[u] for u in self.order]
        conv = lambda v: 'yes' if v is True else 'no' if v is False else ('' if v is None else v)
        if fmt == 'csv':
            buf = io.StringIO(); w = csv.writer(buf, delimiter=';'); w.writerow(self.headers(cols))
            for r in rows: w.writerow([conv(r.get(c, '')) for c in cols])
            return ('\ufeff' + buf.getvalue()).encode('utf-8')
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Profiles'
        thin = Border(*[Side(style='thin', color='D9D9D9')] * 4)
        ws.append(self.headers(cols))
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(name='Arial', bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='305496')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = thin
            ws.column_dimensions[cell.column_letter].width = sc.WIDTHS.get(c, 13)
        ws.row_dimensions[1].height = 30; ws.freeze_panes = 'A2'
        for r in rows:
            ws.append([None] * len(cols)); n = ws.max_row
            for i, c in enumerate(cols, 1):
                v = conv(r.get(c, '')); cell = ws.cell(row=n, column=i, value=v)
                cell.font = Font(name='Arial', size=10); cell.border = thin
                if isinstance(v, (int, float)): cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal='right')
                elif c == 'bio': cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(cols)).column_letter}{ws.max_row}'
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


JOB = Job()
HTML = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'dashboard.html'), encoding='utf-8').read() \
    if os.path.exists(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'dashboard.html')) else '<h1>dashboard.html missing</h1>'


# ---------------------------------------------------------------- http
class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass

    def _send(self, code, body, ctype='application/json', extra=None):
        if isinstance(body, str): body = body.encode('utf-8')
        self.send_response(code); self.send_header('Content-Type', ctype); self.send_header('Content-Length', str(len(body)))
        for k, v in (extra or {}).items(): self.send_header(k, v)
        self.end_headers(); self.wfile.write(body)

    def _json(self, obj, code=200): self._send(code, json.dumps(obj, ensure_ascii=False))

    def do_GET(self):
        u = urlparse(self.path); qs = parse_qs(u.query)
        if u.path == '/':
            return self._send(200, HTML.replace('__FIELDS__', json.dumps([[k, l] for k, l, _ in FIELD_GROUPS])), 'text/html; charset=utf-8')
        if u.path == '/api/status': return self._json(JOB.status())
        if u.path == '/api/results':
            n = int(qs.get('n', ['60'])[0]); return self._json({'rows': JOB.recent(n), 'nposts': sc.NPOSTS})
        if u.path == '/api/proxies': return self._json(POOL.list())
        if u.path == '/api/proxies/test_status': return self._json(TESTALL.status())
        if u.path == '/api/found': return self._send(200, '\n'.join(sorted(JOB.found)), 'text/plain; charset=utf-8')
        if u.path == '/api/download':
            fmt = qs.get('fmt', ['xlsx'])[0]
            if fmt == 'found':
                return self._send(200, '\n'.join(sorted(JOB.found)) + '\n', 'text/plain; charset=utf-8',
                                  {'Content-Disposition': 'attachment; filename="found_usernames.txt"'})
            fields = [f for f in qs.get('fields', [''])[0].split(',') if f]; nposts = int(qs.get('nposts', ['5'])[0])
            data = JOB.export(fmt, fields, nposts)
            ctype = 'text/csv' if fmt == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            return self._send(200, data, ctype, {'Content-Disposition': f'attachment; filename="ig_profiles.{fmt}"'})
        self._send(404, 'not found', 'text/plain')

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        body = json.loads(self.rfile.read(n) or b'{}') if n else {}
        p = self.path
        if p == '/api/start':
            if not body.get('usernames'): return self._json({'ok': False, 'error': 'No usernames given.'}, 400)
            if not POOL.available(): return self._json({'ok': False, 'error': 'No connection available: enable the direct connection or add a proxy.'}, 400)
            ok = JOB.start(body); return self._json({'ok': ok, 'error': None if ok else 'A run is already in progress.'})
        if p == '/api/stop': JOB.stop.set(); return self._json({'ok': True})
        if p == '/api/clear':
            if JOB.running(): return self._json({'ok': False, 'error': 'Stop the run first.'}, 400)
            JOB.clear(); return self._json({'ok': True})
        if p == '/api/proxies/add':
            if not body.get('url', '').strip(): return self._json({'ok': False, 'error': 'Proxy URL is empty.'}, 400)
            return self._json({'ok': True, 'id': POOL.add(body['url'], body.get('rotate', ''))})
        if p == '/api/proxies/import':
            text = body.get('text', '')
            if not text.strip(): return self._json({'ok': False, 'error': 'File is empty.'}, 400)
            added, skipped = POOL.add_many(text.splitlines())
            return self._json({'ok': True, 'added': added, 'skipped': skipped})
        if p == '/api/proxies/remove': POOL.remove(body['id']); return self._json({'ok': True})
        if p == '/api/proxies/remove_many': return self._json({'ok': True, 'removed': POOL.remove_many(body.get('ids', []))})
        if p == '/api/proxies/toggle': POOL.toggle(body['id'], body['enabled']); return self._json({'ok': True})
        if p == '/api/proxies/test': return self._json(POOL.test(body['id']))
        if p == '/api/proxies/test_many':
            ids = body.get('ids') or [x['id'] for x in POOL.list()]
            job = TESTALL.start(ids)
            return self._json({'ok': job is not None, 'error': None if job is not None else 'A test-all run is already in progress.'})
        if p == '/api/proxies/test_all_stop': TESTALL.stop.set(); return self._json({'ok': True})
        self._json({}, 404)


def main():
    try: import openpyxl  # noqa
    except ImportError:
        import subprocess; print('installing openpyxl...', flush=True)
        subprocess.check_call([sys.executable, '-m', 'pip', 'install', '-q', 'openpyxl'])
    srv = ThreadingHTTPServer((HOST, PORT), H); url = f'http://{HOST}:{PORT}'
    print(f'Dashboard: {url}   (Ctrl+C to quit)', flush=True)
    if '--no-browser' not in sys.argv: threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try: srv.serve_forever()
    except KeyboardInterrupt: pass


if __name__ == '__main__':
    main()
